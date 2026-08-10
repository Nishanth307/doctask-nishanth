from pathlib import Path

from docx import Document

from app.ingestion.loaders.base import DocumentLoader, LoadedDocument


class DocxLoader(DocumentLoader):
    def load(self, file_path: Path) -> LoadedDocument:
        document = Document(str(file_path))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        text = "\n".join(paragraphs)

        return LoadedDocument(
            text=text,
            source=str(file_path),
            metadata={
                "format": "docx",
            },
        )
